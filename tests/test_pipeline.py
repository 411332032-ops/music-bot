"""pipeline.py：處理一張圖的完整流程（分析 → 存照片 → 寫 Excel → 準備回覆內容）。

分析器用假的，檔案與 Excel 用真的（放在 tmp_path）。
"""

from datetime import datetime

import pytest
from openpyxl import load_workbook

from music_bot import pipeline
from music_bot.analyzer import AnalysisFailure
from music_bot.pipeline import MAX_IMAGE_BYTES, Processor
from music_bot.schema import Analysis
from music_bot.storage import ExcelStore

WHEN = datetime(2026, 10, 21, 14, 3, 22)
JPEG = b"\xff\xd8\xff\xe0" + b"0" * 64


def analysis(**kw) -> Analysis:
    base = dict(
        is_music_related=True,
        type="專輯封面",
        title="ドライフラワー",
        album_title="某專輯",
        release_type="EP",
        artist="優里",
        genre_estimate=["J-Pop", "抒情"],
        lyrics_text="",
        confidence="高",
    )
    base.update(kw)
    return Analysis(**base)


class FakeAnalyzer:
    def __init__(self, result):
        self.result = result
        self.calls = []

    async def analyze(self, image, ext, caption=None):
        self.calls.append({"image": image, "ext": ext, "caption": caption})
        if isinstance(self.result, BaseException):
            raise self.result
        return self.result


def make(tmp_path, result=None):
    analyzer = FakeAnalyzer(result if result is not None else analysis())
    store = ExcelStore(tmp_path / "music_records.xlsx", sleep=lambda s: None)
    proc = Processor(analyzer=analyzer, store=store, data_dir=tmp_path, photos_dir=tmp_path / "photos")
    return proc, analyzer


def sheet(tmp_path):
    return load_workbook(tmp_path / "music_records.xlsx").active


def row(tmp_path, n=2):
    return [c.value for c in sheet(tmp_path)[n]]


# ---------- 音樂相關 ----------


async def test_music_image_is_saved_recorded_and_counted(tmp_path):
    proc, _ = make(tmp_path)
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.daily_count == 1
    assert r.outcome.image_type == "專輯封面"
    assert r.outcome.confidence == "高"
    photo = tmp_path / "photos" / "2026-10-21 14-03-22 ドライフラワー.jpg"
    assert photo.read_bytes() == JPEG
    cells = row(tmp_path)
    assert cells[0] == WHEN
    assert cells[1:4] == ["小明", "專輯封面", "ドライフラワー"]
    assert cells[4:6] == ["某專輯", "EP"]
    assert cells[6:8] == ["優里", "J-Pop / 抒情"]
    assert cells[8] in (None, "")  # 沒有歌詞
    assert cells[9] == 1
    assert cells[10] == "photos/2026-10-21 14-03-22 ドライフラワー.jpg"
    assert cells[11] == "高"


async def test_analyzer_gets_the_detected_extension_and_caption(tmp_path):
    proc, fake = make(tmp_path)
    await proc.process(JPEG, "小明", WHEN, caption="某段說明")
    assert fake.calls[0]["ext"] == ".jpg"
    assert fake.calls[0]["caption"] == "某段說明"
    assert fake.calls[0]["image"] == JPEG


async def test_second_music_image_increments_the_count(tmp_path):
    proc, _ = make(tmp_path)
    await proc.process(JPEG, "小明", WHEN)
    r = await proc.process(JPEG, "小華", WHEN)
    assert r.daily_count == 2


async def test_same_second_same_title_photos_do_not_overwrite_each_other(tmp_path):
    proc, _ = make(tmp_path)
    await proc.process(JPEG, "a", WHEN)
    await proc.process(JPEG, "b", WHEN)
    names = sorted(p.name for p in (tmp_path / "photos").iterdir())
    assert names == [
        "2026-10-21 14-03-22 ドライフラワー-2.jpg",
        "2026-10-21 14-03-22 ドライフラワー.jpg",
    ]


async def test_backfilled_image_uses_its_original_upload_time_and_day(tmp_path):
    proc, _ = make(tmp_path)
    await proc.process(JPEG, "a", datetime(2026, 10, 22, 9, 0, 0))
    r = await proc.process(JPEG, "b", datetime(2026, 10, 21, 8, 0, 0))  # 補處理昨天的
    assert r.daily_count == 1
    assert row(tmp_path, 3)[0] == datetime(2026, 10, 21, 8, 0, 0)


# ---------- 非音樂 ----------


async def test_non_music_is_recorded_with_description_and_does_not_count(tmp_path):
    proc, _ = make(tmp_path)
    await proc.process(JPEG, "小明", WHEN)
    drink = analysis(
        is_music_related=False,
        type="非音樂相關",
        title="手搖飲料",
        artist="無",
        genre_estimate=[],
        lyrics_text="一杯珍珠奶茶，杯身有品牌標誌",
        confidence="高",
    )
    proc.analyzer.result = drink
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.daily_count == 1
    assert r.outcome.image_type == "非音樂相關"
    assert r.outcome.confidence == "無"
    cells = row(tmp_path, 3)
    assert cells[2] == "非音樂相關"
    assert cells[3] == "手搖飲料"
    assert cells[4:6] == ["無", "無"]  # 專輯名稱、發行類型
    assert cells[6:8] == ["無", "無"]  # 歌手、曲風
    assert cells[8] == "一杯珍珠奶茶，杯身有品牌標誌"
    assert cells[9] == 1
    assert cells[11] == "無"
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 手搖飲料.jpg").exists()


# ---------- 失敗 ----------


async def test_analysis_failure_with_salvaged_title(tmp_path):
    proc, _ = make(tmp_path, AnalysisFailure("格式不符", salvaged_title="一張風景照"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.is_failure
    assert r.outcome.title == "一張風景照"
    assert r.outcome.failure_reason == "格式不符"
    assert r.daily_count == 0
    cells = row(tmp_path)
    assert cells[2] == "處理失敗"
    assert cells[3] == "一張風景照"
    assert cells[4:6] == ["無", "無"]
    assert cells[6:8] == ["無", "無"]
    assert cells[8] == "格式不符"
    assert cells[9] == 0
    assert cells[11] == "無"
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 一張風景照.jpg").exists()  # 原圖照存


async def test_analysis_failure_without_content_uses_fixed_text(tmp_path):
    proc, _ = make(tmp_path, AnalysisFailure("逾時"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.title == "處理失敗，無法判斷圖片內容"
    assert r.outcome.lyrics == "逾時"
    assert row(tmp_path)[3] == "處理失敗，無法判斷圖片內容"
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 處理失敗.jpg").exists()


async def test_failure_does_not_increase_a_running_count(tmp_path):
    proc, fake = make(tmp_path)
    await proc.process(JPEG, "a", WHEN)
    fake.result = AnalysisFailure("逾時")
    r = await proc.process(JPEG, "b", WHEN)
    assert r.daily_count == 1


async def test_unknown_format_skips_the_analyzer_and_keeps_the_file_as_bin(tmp_path):
    proc, fake = make(tmp_path)
    r = await proc.process(b"this is plain text, not an image", "小明", WHEN)
    assert fake.calls == []
    assert r.outcome.failure_reason == "圖片損壞或格式不支援"
    assert [p.suffix for p in (tmp_path / "photos").iterdir()] == [".bin"]


async def test_oversized_image_is_rejected_without_saving(tmp_path):
    proc, fake = make(tmp_path)
    r = await proc.process(JPEG + b"0" * MAX_IMAGE_BYTES, "小明", WHEN)
    assert fake.calls == []
    assert r.outcome.failure_reason == "圖片過大"
    assert not (tmp_path / "photos").exists() or list((tmp_path / "photos").iterdir()) == []
    assert row(tmp_path)[10] == "無"


async def test_download_failure_is_recorded_without_a_photo(tmp_path):
    proc, fake = make(tmp_path)
    r = await proc.process(None, "小明", WHEN)
    assert fake.calls == []
    assert r.outcome.failure_reason == "下載失敗"
    assert row(tmp_path)[2] == "處理失敗"
    assert row(tmp_path)[10] == "無"


async def test_photo_save_failure_turns_the_result_into_a_failure_row(monkeypatch, tmp_path):
    def boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(pipeline.storage, "save_photo", boom)
    proc, _ = make(tmp_path)
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.is_failure
    assert r.outcome.failure_reason == "內部錯誤"
    assert r.daily_count == 0
    assert row(tmp_path)[10] == "無"


# ---------- 防護 3、4 串起來的效果 ----------


async def test_evil_title_cannot_move_the_photo_outside_photos_dir(tmp_path):
    proc, _ = make(tmp_path, analysis(title="..\\..\\evil:*?"))
    await proc.process(JPEG, "小明", WHEN)
    files = [p for p in tmp_path.rglob("*") if p.is_file() and p.suffix == ".jpg"]
    assert len(files) == 1
    assert files[0].parent == tmp_path / "photos"


async def test_formula_title_artist_and_uploader_are_stored_as_text(tmp_path):
    proc, _ = make(tmp_path, analysis(title="=1+1", artist="@SUM(1,2)"))
    await proc.process(JPEG, "+cmd", WHEN)
    ws = sheet(tmp_path)
    for col in (2, 4, 7):  # 上傳者、歌曲名稱、歌手
        cell = ws.cell(row=2, column=col)
        assert cell.data_type != "f"
        assert cell.value.startswith("'")


async def test_formula_in_a_non_music_description_is_stored_as_text(tmp_path):
    """歌詞/文字欄對非音樂圖片存的是內容說明，模型給的文字一樣要過防護 4。"""
    proc, _ = make(
        tmp_path,
        analysis(
            is_music_related=False,
            type="非音樂相關",
            title="=1+1",
            artist="無",
            genre_estimate=[],
            lyrics_text="-2+3",
        ),
    )
    await proc.process(JPEG, "小明", WHEN)
    ws = sheet(tmp_path)
    for col in (4, 9):  # 歌曲名稱欄（放物品說明）、歌詞/文字欄（放內容說明）
        cell = ws.cell(row=2, column=col)
        assert cell.data_type != "f"
        assert cell.value.startswith("'")


async def test_known_failure_reason_is_recorded_without_analysis_or_photo(tmp_path):
    """附件太大時 bot 根本不下載，直接告訴 Processor 原因。"""
    proc, fake = make(tmp_path)
    r = await proc.process(None, "小明", WHEN, failure="圖片過大")
    assert fake.calls == []
    assert r.outcome.failure_reason == "圖片過大"
    assert row(tmp_path)[2] == "處理失敗"
    assert row(tmp_path)[8] == "圖片過大"
    assert row(tmp_path)[10] == "無"


# ---------- 歌詞轉錄功能已停用 ----------


async def test_music_image_never_stores_lyrics_even_if_the_model_returns_text(tmp_path):
    proc, _ = make(tmp_path, analysis(lyrics_text="模型自己多給的一段文字\n第二行"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyrics == ""
    assert row(tmp_path)[8] in (None, "")  # Excel 的歌詞/文字欄是空的


@pytest.mark.parametrize("kind", ["專輯封面", "歌詞手稿", "其他音樂相關"])
async def test_every_music_type_has_its_lyrics_cleared(tmp_path, kind):
    proc, _ = make(tmp_path, analysis(type=kind, lyrics_text="不該被保存的文字"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyrics == ""


async def test_non_music_description_is_still_kept(tmp_path):
    """非音樂的「歌詞/文字」欄放的是內容說明，不是歌詞，所以照樣保留。"""
    proc, _ = make(
        tmp_path,
        analysis(
            is_music_related=False,
            type="非音樂相關",
            title="手搖飲料",
            artist="無",
            genre_estimate=[],
            lyrics_text="一杯珍珠奶茶",
        ),
    )
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyrics == "一杯珍珠奶茶"


# ---------- 專輯名稱與發行類型（專輯／EP／單曲） ----------


async def test_album_and_release_type_reach_the_outcome_and_excel(tmp_path):
    proc, _ = make(tmp_path, analysis(album_title="嗨！人間 Hi! world", release_type="EP"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.album_title == "嗨！人間 Hi! world"
    assert r.outcome.release_type == "EP"
    cells = row(tmp_path)
    assert cells[4:6] == ["嗨！人間 Hi! world", "EP"]


async def test_unknown_album_and_release_type_are_written_as_unknown(tmp_path):
    proc, _ = make(tmp_path, analysis(album_title="不明", release_type="不明"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert (r.outcome.album_title, r.outcome.release_type) == ("不明", "不明")
    assert row(tmp_path)[4:6] == ["不明", "不明"]


async def test_non_music_always_gets_none_word_for_album_and_release_type(tmp_path):
    """模型就算在非音樂圖片上亂填專輯名稱，也不採用。"""
    drink = analysis(
        is_music_related=False,
        type="非音樂相關",
        title="手搖飲料",
        artist="無",
        genre_estimate=[],
        album_title="某專輯",
        release_type="EP",
    )
    proc, _ = make(tmp_path, drink)
    r = await proc.process(JPEG, "小明", WHEN)
    assert (r.outcome.album_title, r.outcome.release_type) == ("無", "無")
    assert row(tmp_path)[4:6] == ["無", "無"]


async def test_music_image_never_gets_none_word_for_release_type(tmp_path):
    proc, _ = make(tmp_path, analysis(release_type="無"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.release_type == "不明"


async def test_failure_rows_have_none_word_for_album_and_release_type(tmp_path):
    proc, _ = make(tmp_path, AnalysisFailure("逾時"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert (r.outcome.album_title, r.outcome.release_type) == ("無", "無")
    assert row(tmp_path)[4:6] == ["無", "無"]


async def test_photo_is_named_after_the_song_when_known(tmp_path):
    proc, _ = make(tmp_path, analysis(title="某首歌", album_title="某專輯"))
    await proc.process(JPEG, "小明", WHEN)
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 某首歌.jpg").exists()


async def test_photo_is_named_after_the_album_when_the_song_is_unknown(tmp_path):
    proc, _ = make(tmp_path, analysis(title="不明", album_title="嗨！人間"))
    await proc.process(JPEG, "小明", WHEN)
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 嗨！人間.jpg").exists()


async def test_photo_name_falls_back_to_unknown_when_both_names_are_unknown(tmp_path):
    proc, _ = make(tmp_path, analysis(title="不明", album_title="不明"))
    await proc.process(JPEG, "小明", WHEN)
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 不明.jpg").exists()


async def test_photo_is_named_after_the_album_when_the_song_line_is_a_two_names_note(tmp_path):
    proc, _ = make(tmp_path, analysis(title="不明（圖上另有：嗨！人間）", album_title="小深情"))
    await proc.process(JPEG, "小明", WHEN)
    assert (tmp_path / "photos" / "2026-10-21 14-03-22 小深情.jpg").exists()


async def test_two_names_note_is_stored_in_the_song_column(tmp_path):
    proc, _ = make(tmp_path, analysis(title="不明（圖上另有：嗨！人間）", album_title="小深情"))
    await proc.process(JPEG, "小明", WHEN)
    cells = row(tmp_path)
    assert cells[3] == "不明（圖上另有：嗨！人間）" and cells[4] == "小深情"


# ---------- 歌名不明的歌詞截圖：留下一小句搜尋用的關鍵句（只給回覆做連結，不存 Excel） ----------

SNIPPET = "測試用的關鍵句子"


def all_cell_text(tmp_path) -> str:
    return " | ".join(str(c.value) for r in sheet(tmp_path).iter_rows() for c in r if c.value is not None)


async def test_snippet_is_kept_for_a_lyrics_photo_with_unknown_song(tmp_path):
    proc, _ = make(tmp_path, analysis(type="歌詞手稿", title="不明", album_title="不明", artist="無", lyric_snippet=SNIPPET))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyric_snippet == SNIPPET


async def test_snippet_is_never_written_to_excel(tmp_path):
    proc, _ = make(tmp_path, analysis(type="歌詞手稿", title="不明", album_title="不明", artist="無", lyric_snippet=SNIPPET))
    await proc.process(JPEG, "小明", WHEN)
    assert SNIPPET not in all_cell_text(tmp_path)


async def test_snippet_is_dropped_when_the_song_is_known(tmp_path):
    proc, _ = make(tmp_path, analysis(type="歌詞手稿", title="某首歌", lyric_snippet=SNIPPET))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyric_snippet == ""


@pytest.mark.parametrize("kind", ["專輯封面", "其他音樂相關"])
async def test_snippet_is_dropped_for_other_music_types(tmp_path, kind):
    proc, _ = make(tmp_path, analysis(type=kind, title="不明", lyric_snippet=SNIPPET))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyric_snippet == ""


async def test_snippet_is_dropped_for_non_music(tmp_path):
    drink = analysis(
        is_music_related=False,
        type="非音樂相關",
        title="手搖飲料",
        artist="無",
        genre_estimate=[],
        lyric_snippet=SNIPPET,
    )
    proc, _ = make(tmp_path, drink)
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyric_snippet == ""


async def test_failure_outcome_has_no_snippet(tmp_path):
    proc, _ = make(tmp_path, AnalysisFailure("逾時"))
    r = await proc.process(JPEG, "小明", WHEN)
    assert r.outcome.lyric_snippet == ""


# ---------- log：處理過程每一步都要看得到（記分析結果，但不記上傳者暱稱） ----------


async def test_processing_steps_are_logged_without_the_uploader_name(tmp_path, caplog):
    import logging

    proc, _ = make(tmp_path)
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(JPEG, "小明", WHEN)
    text = "\n".join(caplog.messages)
    assert "開始處理" in text and f"{len(JPEG)} bytes" in text
    assert "分析完成" in text and "專輯封面" in text and "高" in text
    assert "照片已存檔" in text
    assert "Excel 已寫入" in text and "當日音樂數=1" in text
    assert "小明" not in text  # 上傳者暱稱不進 log


async def test_analysis_failure_is_logged_with_its_reason(tmp_path, caplog):
    import logging

    proc, _ = make(tmp_path, AnalysisFailure("逾時"))
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(JPEG, "小明", WHEN)
    assert any("分析失敗" in m and "逾時" in m for m in caplog.messages)
    assert any("Excel 已寫入" in m and "處理失敗" in m for m in caplog.messages)


async def test_image_that_was_never_downloaded_is_logged(tmp_path, caplog):
    import logging

    proc, _ = make(tmp_path)
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(None, "小明", WHEN, failure="下載失敗")
    assert any("沒有取得圖片" in m and "下載失敗" in m for m in caplog.messages)


async def test_analysis_result_is_logged(tmp_path, caplog):
    import logging

    proc, _ = make(tmp_path)
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(JPEG, "小明", WHEN)
    line = next(m for m in caplog.messages if m.startswith("分析結果"))
    for expected in ("ドライフラワー", "某專輯", "EP", "優里", "J-Pop / 抒情"):
        assert expected in line


async def test_non_music_description_is_in_the_result_log_but_shortened_and_on_one_line(tmp_path, caplog):
    import logging

    long_text = "第一行" + chr(10) + "第二行" + "很長" * 200
    proc, _ = make(tmp_path, analysis(is_music_related=False, type="非音樂相關", title="一碗拉麵", lyrics_text=long_text))
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(JPEG, "小明", WHEN)
    line = next(m for m in caplog.messages if m.startswith("分析結果"))
    assert "一碗拉麵" in line and "第一行" in line
    assert chr(10) not in line  # 換行要變成跳脫字元，不能讓一筆 log 斷成好幾行或偽造另一行
    assert len(line) < 400


async def test_lyric_snippet_is_not_written_to_the_log(tmp_path, caplog):
    import logging

    proc, _ = make(tmp_path, analysis(type="歌詞手稿", title="不明", lyric_snippet="這是一句只拿來搜尋的歌詞"))
    with caplog.at_level(logging.INFO, logger="music_bot.pipeline"):
        await proc.process(JPEG, "小明", WHEN)
    assert not any("只拿來搜尋" in m for m in caplog.messages)


# ---------- 上傳者 ID 與個人當日數 ----------


async def test_uploader_id_is_recorded_and_the_user_count_is_returned(tmp_path):
    proc, _ = make(tmp_path)
    r1 = await proc.process(JPEG, "小明", WHEN, uploader_id="111")
    r2 = await proc.process(JPEG, "小華", WHEN, uploader_id="222")
    r3 = await proc.process(JPEG, "小明", WHEN, uploader_id="111")
    assert [r.user_count for r in (r1, r2, r3)] == [1, 1, 2]
    assert [row(tmp_path, n)[12] for n in (2, 3, 4)] == ["111", "222", "111"]


async def test_without_an_uploader_id_there_is_no_user_count(tmp_path):
    proc, _ = make(tmp_path)
    assert (await proc.process(JPEG, "小明", WHEN)).user_count is None


async def test_processor_stats_reads_the_excel(tmp_path):
    proc, _ = make(tmp_path)
    await proc.process(JPEG, "小明", WHEN, uploader_id="111")
    st = await proc.stats("111", WHEN.date())
    assert st.total_music == 1 and st.user_today == 1 and st.top_artists == [("優里", 1)]
